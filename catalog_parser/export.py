import csv
from collections import defaultdict
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from catalog_parser.models import Product

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="D0D7E2")
PRICE_FORMAT = '£#,##0.00'

# (заголовок, ширина, функция получения значения)
COLUMNS = [
    ("Название", 45, lambda p: p.title),
    ("Категория", 18, lambda p: p.category),
    ("Цена", 11, lambda p: p.price),
    ("Цена без налога", 11, lambda p: p.price_excl_tax),
    ("В наличии, шт", 10, lambda p: p.in_stock),
    ("Рейтинг", 10, lambda p: "★" * p.rating + "☆" * (5 - p.rating)),
    ("Отзывы", 9, lambda p: p.reviews),
    ("Артикул (UPC)", 19, lambda p: p.upc),
    ("Описание", 70, lambda p: p.description[:1000]),
    ("Фото", 12, lambda p: p.image_url),
]
PRICE_COLUMNS = {"Цена", "Цена без налога"}


def _style_header(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _products_sheet(ws, products: list[Product]) -> None:
    ws.title = "Товары"
    ws.append([name for name, _, _ in COLUMNS])
    for p in products:
        ws.append([get(p) for _, _, get in COLUMNS])

    _style_header(ws, [width for _, width, _ in COLUMNS])
    names = [name for name, _, _ in COLUMNS]
    title_col, photo_col = names.index("Название") + 1, names.index("Фото") + 1
    border = Border(bottom=THIN)
    for row, p in enumerate(products, start=2):
        for col, name in enumerate(names, start=1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=name in {"Название", "Описание"})
            if name in PRICE_COLUMNS:
                cell.number_format = PRICE_FORMAT
        # Название — ссылка на карточку товара, фото — ссылка на картинку
        ws.cell(row=row, column=title_col).hyperlink = p.url
        ws.cell(row=row, column=title_col).font = LINK_FONT
        if p.image_url:
            photo = ws.cell(row=row, column=photo_col, value="открыть")
            photo.hyperlink, photo.font = p.image_url, LINK_FONT
        ws.row_dimensions[row].height = 45
    ws.auto_filter.ref = ws.dimensions


def category_summary(products: list[Product]) -> list[tuple[str, int, float, float, float, int]]:
    """(категория, товаров, средняя цена, мин, макс, всего на складе), по убыванию числа товаров."""
    groups: dict[str, list[Product]] = defaultdict(list)
    for p in products:
        groups[p.category].append(p)
    rows = [
        (
            cat,
            len(items),
            round(sum(p.price for p in items) / len(items), 2),
            min(p.price for p in items),
            max(p.price for p in items),
            sum(p.in_stock for p in items),
        )
        for cat, items in groups.items()
    ]
    return sorted(rows, key=lambda r: (-r[1], r[0]))


def _summary_sheet(wb: Workbook, products: list[Product]) -> None:
    ws = wb.create_sheet("Сводка")
    headers = ["Категория", "Товаров", "Средняя цена", "Мин. цена", "Макс. цена", "На складе, шт"]
    ws.append(headers)
    rows = category_summary(products)
    for r in rows:
        ws.append(list(r))
    total_row = len(rows) + 2
    last = total_row - 1
    ws.append([
        "ИТОГО",
        f"=SUM(B2:B{last})",
        f"=ROUND(SUMPRODUCT(B2:B{last},C2:C{last})/B{total_row},2)",
        f"=MIN(D2:D{last})",
        f"=MAX(E2:E{last})",
        f"=SUM(F2:F{last})",
    ])
    _style_header(ws, [24, 10, 14, 12, 12, 14])
    for row in ws.iter_rows(min_row=2, max_row=total_row, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = PRICE_FORMAT
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.border = Border(top=Side(style="medium", color="1E3A5F"))

    if rows:
        top = min(len(rows), 10)
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Топ категорий по количеству товаров"
        chart.legend = None
        chart.varyColors = False
        chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=top + 1))
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=top + 1))
        chart.series[0].graphicalProperties.solidFill = "1E3A5F"
        chart.series[0].graphicalProperties.line.noFill = True
        chart.gapWidth = 60
        # openpyxl 3.1 по умолчанию помечает оси удалёнными — Excel рисует диаграмму без подписей категорий
        chart.x_axis.delete = False
        # Шкалу значений не показываем: числа подписаны прямо у столбцов
        chart.y_axis.delete = True
        chart.y_axis.majorGridlines = None
        # Всё, что не задано явно, Excel показывает: имя ряда, категорию, маркер
        chart.dataLabels = DataLabelList(
            showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False
        )
        chart.title.overlay = False
        chart.x_axis.scaling.orientation = "maxMin"  # самая большая категория сверху
        chart.height, chart.width = 9, 16
        ws.add_chart(chart, "H2")


def _info_sheet(wb: Workbook, products: list[Product], source: str, failed: int, seconds: float) -> None:
    ws = wb.create_sheet("Инфо")
    for row in [
        ("Источник", source),
        ("Дата выгрузки", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Товаров собрано", len(products)),
        ("Не удалось загрузить", failed),
        ("Время работы, с", round(seconds, 1)),
    ]:
        ws.append(row)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 45
    for cell in ws["A"]:
        cell.font = Font(bold=True)


def to_excel(
    products: list[Product], path: Path, source: str = "", failed: int = 0, seconds: float = 0.0
) -> None:
    wb = Workbook()
    _products_sheet(wb.active, products)
    _summary_sheet(wb, products)
    _info_sheet(wb, products, source, failed, seconds)
    wb.save(path)


def to_csv(products: list[Product], path: Path) -> None:
    # utf-8-sig — чтобы Excel открыл кириллицу без кракозябр
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(Product.__dataclass_fields__), delimiter=";")
        writer.writeheader()
        writer.writerows(asdict(p) for p in products)
