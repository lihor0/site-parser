import csv

from openpyxl import load_workbook

from catalog_parser.export import category_summary, to_csv, to_excel
from catalog_parser.models import Product


def make(title: str, category: str, price: float, stock: int = 1, rating: int = 3) -> Product:
    return Product(
        title=title, category=category, price=price, price_excl_tax=price, in_stock=stock,
        rating=rating, reviews=0, upc="u-" + title, description="Описание " + title,
        url=f"https://shop.example/{title}", image_url=f"https://shop.example/{title}.jpg",
    )


PRODUCTS = [
    make("A", "Mystery", 10.0, stock=5),
    make("B", "Mystery", 30.0, stock=1, rating=5),
    make("C", "Travel", 20.0, stock=2, rating=0),
]


def test_summary():
    assert category_summary(PRODUCTS) == [
        ("Mystery", 2, 20.0, 10.0, 30.0, 6),
        ("Travel", 1, 20.0, 20.0, 20.0, 2),
    ]


def test_excel(tmp_path):
    path = tmp_path / "out.xlsx"
    to_excel(PRODUCTS, path, source="https://shop.example", failed=1, seconds=2.5)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Товары", "Сводка", "Инфо"]

    ws = wb["Товары"]
    assert ws["A1"].value == "Название"
    assert ws.max_row == 4
    assert ws["A2"].value == "A"
    assert ws["A2"].hyperlink.target == "https://shop.example/A"
    assert ws["C3"].value == 30.0
    assert ws["C3"].number_format.startswith("£")
    assert ws["F3"].value == "★★★★★"
    assert ws["F4"].value == "☆☆☆☆☆"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:J4"

    summary = wb["Сводка"]
    assert summary["A2"].value == "Mystery"
    assert summary["A4"].value == "ИТОГО"
    assert summary["B4"].value == "=SUM(B2:B3)"
    assert len(summary._charts) == 1
    chart = summary._charts[0]
    assert chart.x_axis.delete is False  # иначе Excel не покажет названия категорий
    assert chart.dataLabels.showVal

    info = {row[0].value: row[1].value for row in wb["Инфо"].iter_rows()}
    assert info["Товаров собрано"] == 3
    assert info["Не удалось загрузить"] == 1


def test_csv_opens_in_excel(tmp_path):
    path = tmp_path / "out.csv"
    to_csv(PRODUCTS, path)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM для Excel
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0]["title"] == "A"
    assert rows[2]["description"] == "Описание C"
